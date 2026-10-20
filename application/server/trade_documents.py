# -*- coding: utf-8 -*-
"""本机贸易单据读取与字段候选提取。

解析器先提取文本和结构化单元格，再按字段语义匹配；页码、段落、工作表和
单元格仅作为证据位置，不参与字段识别规则。
"""
from __future__ import annotations

import io
import base64
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import Response

router = APIRouter(prefix="/api/documents", tags=["trade-documents"])
from server.document_sources import router as source_router
router.include_router(source_router)

# Spreadsheet edit sessions are local, short-lived copies. The uploaded source
# is never overwritten and sessions expire after one hour.
_WORKBOOK_SESSIONS: dict[str, dict[str, Any]] = {}
_WORKBOOK_SESSION_TTL = 60 * 60
_WORKBOOK_MAX_CELLS = 100_000


def _get_workbook_session(session_id: str) -> dict[str, Any]:
    session = _WORKBOOK_SESSIONS.get(session_id)
    if not session or time.time() - session["created"] > _WORKBOOK_SESSION_TTL:
        _WORKBOOK_SESSIONS.pop(session_id, None)
        raise HTTPException(404, "工作簿编辑会话已过期，请从冲突来源重新打开。")
    return session


def _excel_color(color: Any) -> str | None:
    if not color or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    match = re.fullmatch(r"(?:[0-9A-Fa-f]{2})?([0-9A-Fa-f]{6})", color.rgb)
    return f"#{match.group(1)}" if match else None


def _cell_style(cell: Any) -> dict[str, Any]:
    font = cell.font
    fill = cell.fill
    alignment = cell.alignment
    style: dict[str, Any] = {
        "font": {key: value for key, value in {
            "name": font.name, "size": font.sz, "bold": font.bold, "italic": font.italic,
            "underline": bool(font.underline), "strike": font.strike,
            "color": _excel_color(font.color),
        }.items() if value is not None and value is not False},
        "fill": _excel_color(fill.fgColor) if fill.patternType == "solid" else None,
        "alignment": {key: value for key, value in {
            "horizontal": alignment.horizontal, "vertical": alignment.vertical,
            "wrap": alignment.wrap_text, "rotation": alignment.textRotation,
            "indent": alignment.indent,
        }.items() if value is not None and value is not False},
        "borders": {},
    }
    for side_name in ("left", "right", "top", "bottom"):
        side = getattr(cell.border, side_name)
        if side.style:
            style["borders"][side_name] = {"style": side.style, "color": _excel_color(side.color)}
    return style


def _workbook_payload(content: bytes, focus_sheet: str = "", focus_cell: str = "") -> dict[str, Any]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), data_only=False)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"无法打开 Excel 工作簿：{exc}") from exc
    sheets = []
    styles: list[dict[str, Any]] = []
    style_ids: dict[int, int] = {}
    total_cells = 0
    for sheet in book.worksheets:
        if sheet.max_row * sheet.max_column > _WORKBOOK_MAX_CELLS:
            book.close()
            raise HTTPException(400, f"工作表「{sheet.title}」内容过大，暂不支持网页编辑（最多 {_WORKBOOK_MAX_CELLS} 个单元格）。")
        total_cells += sheet.max_row * sheet.max_column
        if total_cells > _WORKBOOK_MAX_CELLS:
            book.close()
            raise HTTPException(400, f"工作簿内容过大，暂不支持网页编辑（最多 {_WORKBOOK_MAX_CELLS} 个单元格）。")
        rows = []
        row_styles = []
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            rows.append([cell.value.isoformat() if hasattr(cell.value, "isoformat") else
                         cell.value if cell.value is None or isinstance(cell.value, (str, int, float, bool)) else str(cell.value)
                         for cell in row])
            ids = []
            for cell in row:
                style_id = cell.style_id
                if style_id not in style_ids:
                    style_ids[style_id] = len(styles)
                    styles.append(_cell_style(cell))
                ids.append(style_ids[style_id])
            row_styles.append(ids)
        columns = []
        for letter, dimension in sheet.column_dimensions.items():
            columns.append({"letter": letter, "min": dimension.min, "max": dimension.max,
                            "width": dimension.width, "hidden": dimension.hidden})
        row_dimensions = {str(index): {"height": dimension.height, "hidden": dimension.hidden}
                          for index, dimension in sheet.row_dimensions.items()
                          if dimension.height is not None or dimension.hidden}
        merged = []
        for cell_range in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = cell_range.bounds
            merged.append({"minRow": min_row, "maxRow": max_row,
                           "minColumn": min_col, "maxColumn": max_col})
        sheets.append({"name": sheet.title, "rows": rows, "styleIds": row_styles,
                       "maxRow": sheet.max_row, "maxColumn": sheet.max_column,
                       "columns": columns, "rowDimensions": row_dimensions,
                       "defaultColumnWidth": sheet.sheet_format.defaultColWidth or 8.43,
                       "merged": merged})
    book.close()
    if not sheets:
        raise HTTPException(400, "工作簿没有可编辑的工作表。")
    if focus_sheet not in {sheet["name"] for sheet in sheets}:
        focus_sheet = sheets[0]["name"]
    return {"sheets": sheets, "styles": styles, "focusSheet": focus_sheet, "focusCell": focus_cell}


def _xls_workbook_payload(content: bytes, focus_sheet: str = "", focus_cell: str = "") -> dict[str, Any]:
    """Read BIFF workbooks in place; never convert or rewrite the source."""
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"无法打开 Excel 工作簿：{exc}") from exc
    try:
        def color(index):
            rgb = book.colour_map.get(index)
            return '#%02x%02x%02x' % rgb if rgb is not None else None

        line_styles = {1: 'thin', 2: 'medium', 3: 'dashed', 4: 'dotted', 5: 'thick', 6: 'double',
                       7: 'hair', 8: 'mediumDashed', 9: 'dashDot', 10: 'mediumDashDot',
                       11: 'dashDotDot', 12: 'mediumDashDotDot', 13: 'slantDashDot'}
        styles = []
        for xf in book.xf_list:
            font = book.font_list[xf.font_index]
            alignment = xf.alignment
            borders = {}
            for side in ('left', 'right', 'top', 'bottom'):
                line = getattr(xf.border, side + '_line_style')
                if line:
                    borders[side] = {'style': line_styles.get(line, 'thin'),
                                     'color': color(getattr(xf.border, side + '_colour_index'))}
            styles.append({
                'font': {'name': font.name, 'size': font.height / 20, 'bold': bool(font.bold),
                         'italic': bool(font.italic), 'underline': bool(font.underlined),
                         'strike': bool(font.struck_out), 'color': color(font.colour_index)},
                'fill': color(xf.background.pattern_colour_index) if xf.background.fill_pattern == 1 else None,
                'alignment': {'horizontal': {1: 'left', 2: 'center', 3: 'right', 4: 'fill', 5: 'justify',
                                             6: 'center', 7: 'distributed'}.get(alignment.hor_align),
                              'vertical': {0: 'top', 1: 'center', 2: 'bottom'}.get(alignment.vert_align),
                              'wrap': bool(alignment.text_wrapped), 'indent': alignment.indent_level},
                'borders': borders,
            })
        sheets = []
        total_cells = 0
        for sheet in book.sheets():
            total_cells += sheet.nrows * sheet.ncols
            if total_cells > _WORKBOOK_MAX_CELLS:
                raise HTTPException(400, f"工作簿内容过大，暂不支持网页查看（最多 {_WORKBOOK_MAX_CELLS} 个单元格）。")
            rows, row_styles = [], []
            for row_index in range(sheet.nrows):
                values, ids = [], []
                for col_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, col_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, '#ERROR!')
                    values.append(value)
                    ids.append(sheet.cell_xf_index(row_index, col_index))
                rows.append(values)
                row_styles.append(ids)
            sheets.append({
                'name': sheet.name, 'rows': rows, 'styleIds': row_styles,
                'maxRow': sheet.nrows, 'maxColumn': sheet.ncols,
                'columns': [{'min': index + 1, 'max': index + 1, 'width': info.width / 256,
                             'hidden': bool(info.hidden)} for index, info in sheet.colinfo_map.items()],
                'rowDimensions': {str(index + 1): {'height': info.height / 20 if not info.has_default_height else None,
                                                  'hidden': bool(info.hidden)}
                                  for index, info in sheet.rowinfo_map.items()},
                'defaultColumnWidth': sheet.defcolwidth or 8.43,
                'merged': [{'minRow': rlo + 1, 'maxRow': rhi, 'minColumn': clo + 1, 'maxColumn': chi}
                           for rlo, rhi, clo, chi in sheet.merged_cells],
            })
        if not sheets:
            raise HTTPException(400, "工作簿没有可查看的工作表。")
        if focus_sheet not in {sheet['name'] for sheet in sheets}:
            focus_sheet = sheets[0]['name']
        return {'sheets': sheets, 'styles': styles, 'focusSheet': focus_sheet, 'focusCell': focus_cell}
    finally:
        book.release_resources()


def _session_workbook_payload(session: dict[str, Any], sheet: str = "", cell: str = ""):
    reader = _xls_workbook_payload if Path(session['filename']).suffix.lower() == '.xls' else _workbook_payload
    return reader(session['content'], sheet, cell)


@router.post("/workbook-sessions")
async def create_workbook_session(file: UploadFile = File(...)):
    filename = file.filename or "工作簿.xlsx"
    if Path(filename).suffix.lower() not in {'.xls', '.xlsx'}:
        raise HTTPException(400, "工作簿来源查看支持 .xls 和 .xlsx 文件。")
    content = await file.read()
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(413, "Excel 文件不能超过 25 MB。")
    _session_workbook_payload({'filename': filename, 'content': content})
    now = time.time()
    for key, session in list(_WORKBOOK_SESSIONS.items()):
        if now - session["created"] > _WORKBOOK_SESSION_TTL:
            _WORKBOOK_SESSIONS.pop(key, None)
    session_id = uuid.uuid4().hex
    _WORKBOOK_SESSIONS[session_id] = {"created": now, "filename": filename, "content": content}
    return {"sessionId": session_id, "fileName": filename}


@router.get("/workbook-sessions/{session_id}")
async def get_workbook_session(session_id: str, sheet: str = "", cell: str = ""):
    session = _get_workbook_session(session_id)
    payload = _session_workbook_payload(session, sheet, cell)
    payload["fileName"] = session["filename"]
    payload['capabilities'] = {'edit': Path(session['filename']).suffix.lower() == '.xlsx', 'downloadOriginal': True}
    return payload


@router.get("/workbook-sessions/{session_id}/original")
async def download_workbook_original(session_id: str):
    session = _get_workbook_session(session_id)
    media_type = ('application/vnd.ms-excel' if Path(session['filename']).suffix.lower() == '.xls'
                  else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    return Response(session['content'], media_type=media_type,
                    headers={'Content-Disposition': f"attachment; filename*=UTF-8''{quote(session['filename'])}"})


@router.post("/workbook-sessions/{session_id}/save")
async def save_workbook_session(session_id: str, request: Request):
    from openpyxl import load_workbook

    session = _get_workbook_session(session_id)
    if Path(session['filename']).suffix.lower() == '.xls':
        raise HTTPException(400, "旧版 .xls 来源目前只读，请下载原件在 Excel 中编辑。")
    body = await request.json()
    changes = body.get("changes", []) if isinstance(body, dict) else []
    if not isinstance(changes, list) or len(changes) > _WORKBOOK_MAX_CELLS:
        raise HTTPException(400, "单次修改数量无效。")
    book = load_workbook(io.BytesIO(session["content"]))
    try:
        for change in changes:
            if not isinstance(change, dict):
                continue
            sheet_name, coordinate = change.get("sheet"), change.get("cell")
            if sheet_name not in book.sheetnames or not isinstance(coordinate, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", coordinate):
                continue
            value = change.get("value", "")
            if not isinstance(value, (str, int, float, bool)) and value is not None:
                value = str(value)
            book[sheet_name][coordinate] = value if value != "" else None
        output = io.BytesIO()
        book.save(output)
    finally:
        book.close()
    filename = session["filename"].replace('"', "")
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@dataclass
class TextRow:
    """文档中的一行内容；cells 保留表格结构供语义表头识别。"""

    text: str
    location: str
    cells: list[tuple[str, str]] = field(default_factory=list)
    ocr_score: float | None = None
    cell_x: dict[str, float] = field(default_factory=dict)


# Product fields observed across multiple files in the local nine-file corpus.
# Reuse includes transaction/batch data; conflicting values still need review.
REUSABLE_PRODUCT_FIELDS = (
    'item_description', 'item_specification', 'hs_code', 'quantity', 'unit_price',
    'amount', 'currency', 'production_date', 'sample_quantity', 'batch_no',
)

FIELD_DEFS: list[dict[str, Any]] = [
    {"id": "contract_no", "label": "合同号", "group": "单据", "aliases": ["合同号", "合同编号", "合同协议号", "协议编号", "contract no", "contract number", "contract agreement no"]},
    {"id": "document_no", "label": "单据编号", "group": "单据", "aliases": ["pi no", "proforma invoice no", "invoice no", "invoice number", "发票号", "发票号码", "商业发票号", "单据号"]},
    {"id": "document_date", "label": "单据日期", "group": "单据", "aliases": ["contract date", "invoice date", "pi date", "合同日期", "发票日期", "日期", "date"]},
    {"id": "seller_name", "label": "卖方 / 发货人", "group": "交易方", "aliases": ["the seller", "seller", "exporter", "shipper", "卖方", "销售方", "境内发货人", "境内发货人名称", "发货人"]},
    {"id": "seller_address", "label": "卖方地址", "group": "交易方", "aliases": ["seller address", "exporter address", "卖方地址", "销售方地址"]},
    {"id": "buyer_name", "label": "买方 / 购方", "group": "交易方", "aliases": ["the buyer", "buyer", "purchaser", "customer", "买方", "购方", "客户", "境外收货人"]},
    {"id": "buyer_address", "label": "买方地址", "group": "交易方", "aliases": ["buyer address", "purchaser address", "买方地址", "购方地址", "客户地址"]},
    {"id": "consignee_name", "label": "收货人", "group": "交易方", "aliases": ["consignee", "consignee name", "收货人", "收货单位"]},
    {"id": "consignee_address", "label": "收货人地址", "group": "交易方", "aliases": ["consignee address", "收货人地址"]},
    {"id": "item_description", "label": "产品名称", "group": "产品", "aliases": ["commodity and specifications", "decription of goods", "description of goods", "commodity", "description", "product name", "goods description", "品名", "商品名称", "货物名称", "货物描述", "品名和规格"]},
    {"id": "item_specification", "label": "规格型号", "group": "产品", "aliases": ["specification", "specifications", "model", "规格", "规格型号", "包装规格"]},
    {"id": "hs_code", "label": "HS 编码", "group": "产品", "aliases": ["h.s. code", "h.s code", "hs code", "commodity code", "商品编号", "商品编码", "海关编码"]},
    {"id": "quantity", "label": "数量", "group": "产品", "aliases": ["quantity", "qty", "件数", "数量", "成交数量"]},
    {"id": "unit_price", "label": "单价", "group": "产品", "aliases": ["unit price", "price per", "单价"]},
    {"id": "amount", "label": "金额", "group": "产品", "aliases": ["total amount", "amount", "总价", "总金额", "金额", "成交总价"]},
    {"id": "currency", "label": "币种", "group": "产品", "aliases": ["currency", "币种", "币制"]},
    {"id": "trade_term", "label": "贸易条款", "group": "运输", "aliases": ["price term", "price terms", "incoterms", "trade term", "terms of trade", "价格条款", "成交方式", "贸易术语"]},
    {"id": "loading_port", "label": "装货港", "group": "运输", "aliases": ["port of loading", "loading port", "装运港", "装货港", "离境口岸"]},
    {"id": "destination_port", "label": "目的港", "group": "运输", "aliases": ["port of discharge", "destination port", "port of destination", "指运港", "目的港"]},
    {"id": "destination_country", "label": "目的国家/地区", "group": "运输", "aliases": ["final destination", "destination country", "运抵国", "最终目的国", "目的地"]},
    {"id": "shipment_date", "label": "装运时间", "group": "运输", "aliases": ["time of shipment", "shipment date", "etd", "装运期", "装运日期", "出口日期"]},
    {"id": "delivery_date", "label": "交货时间", "group": "运输", "aliases": ["delivery time", "delivery date", "交货期", "交货日期", "交货时间"]},
    {"id": "payment_terms", "label": "付款条款", "group": "付款", "aliases": ["terms of payment", "payment terms", "付款方式", "付款条款"]},
    {"id": "package_type", "label": "包装方式", "group": "包装", "aliases": ["packing list 装箱单", "packing list", "packing", "package type", "包装种类", "包装方式", "包装"]},
    {"id": "package_count", "label": "包装件数", "group": "包装", "aliases": ["package count", "number of packages", "包装件数", "箱数"]},
    {"id": "gross_weight", "label": "毛重", "group": "包装", "aliases": ["gross weight", "gross wt", "毛重"]},
    {"id": "net_weight", "label": "净重", "group": "包装", "aliases": ["net weight", "net wt", "净重"]},
    {"id": "container_count", "label": "柜型与柜量", "group": "运输", "aliases": ["quantity of container", "container count", "container type", "箱型箱量", "柜量", "集装箱"]},
    {"id": "bill_of_lading", "label": "提单号", "group": "运输", "aliases": ["bill of lading", "b/l no", "bl no", "提单号", "提运单号"]},
]

# Additional labels found in pickup notices, inspection reports and customs forms.
_EXTRA_ALIASES = {
    'contract_no': ['进口合同号', '出口合同号'],
    'document_no': ['检验号', '报告编号', '电子底账数据号'],
    'document_date': ['申报日期', '申请日期'],
    'seller_name': ['生产销售单位', '经营单位名称', 'beneficiary', '申请单位（加盖公章）'],
    'buyer_name': ['to'],
    'seller_address': ['beneficiary address'],
    'item_description': ['产品名称', '品名（中文）', '货物名称(中/外文)', '商品名称及规格型号'],
    'quantity': ["q'ty", '申报数量', '数 量', '批量', '数/重量', '数量及单位'],
    'package_count': ['total ctns'],
    'loading_port': ['起始港'],
    'destination_port': ['卸货港'],
    'destination_country': ['destination', '输往国家(地区)'],
    'shipment_date': ['proforma etd', '开航时间'],
    'unit_price': ['申报单价'],
    'amount': ['申报总价', '货物总值'],
    'hs_code': ['商品代码', 'H.S.编码'],
    'package_type': ['包装种类及数量'],
    'container_count': ["q'ty of contianer", "q'ty of container", '箱型/箱量', '集装箱规格、数量及号码'],
    'bill_of_lading': ['提单号为', '提单号', '提单号码'],
}
for _definition in FIELD_DEFS:
    _definition['aliases'].extend(_EXTRA_ALIASES.get(_definition['id'], []))
FIELD_DEFS.extend([
    {'id': 'pickup_no', 'label': '提箱单号', 'group': '运输', 'aliases': ['提箱单号', '提箱号', '提箱凭证号']},
    {'id': 'vessel_name', 'label': '船名', 'group': '运输', 'aliases': ['船名', '中文船名', 'm.v.', 'vessel', '运输工具名称及航次号']},
    {'id': 'voyage_no', 'label': '航次', 'group': '运输', 'aliases': ['航次', 'voyage', 'v.']},
    {'id': 'brand', 'label': '品牌', 'group': '产品', 'aliases': ['品牌']},
    {'id': 'brand_type', 'label': '品牌类型', 'group': '产品', 'aliases': ['品牌类型']},
    {'id': 'processing_method', 'label': '制作或保存方法', 'group': '产品', 'aliases': ['制作或保存方法（如冷冻等）', '制作或保存方法（未冷冻等）', '制作或保存方法']},
    {'id': 'production_date', 'label': '生产日期', 'group': '产品', 'aliases': ['生产日期']},
    {'id': 'sample_quantity', 'label': '样品数量', 'group': '产品', 'aliases': ['样品数量']},
    {'id': 'batch_no', 'label': '生产批号', 'group': '产品', 'aliases': ['批号']},
    {'id': 'cutoff_date', 'label': '截单时间', 'group': '运输', 'aliases': ['截单时间']},
    {'id': 'discharge_port', 'label': '卸货港', 'group': '运输', 'aliases': ['卸货港']},
    {'id': 'declared_weight', 'label': '申报重量（未注明毛/净重）', 'group': '包装', 'aliases': ['申报重量']},
    {'id': 'export_preference', 'label': '出口享惠情况', 'group': '产品', 'aliases': ['出口享惠情况']},
    {'id': 'composition', 'label': '成分含量', 'group': '产品', 'aliases': ['成分含量']},
    {'id': 'quantity_per_container', 'label': '每柜数量', 'group': '产品', 'aliases': ["Q'TY /20GP CTNS", 'quantity per container']},
])

# A discharge port and final destination can both be printed on one notice.
next(f for f in FIELD_DEFS if f['id'] == 'destination_port')['aliases'] = [
    a for a in next(f for f in FIELD_DEFS if f['id'] == 'destination_port')['aliases'] if a != '卸货港']

_FIELDS = {item["id"]: item for item in FIELD_DEFS}
_NORMALIZE = re.compile(r"[^a-z0-9\u3400-\u9fff]+", re.I)
_ALIAS_LOOKUP = sorted(
    ((re.sub(_NORMALIZE, "", alias.lower()), item["id"], alias) for item in FIELD_DEFS for alias in item["aliases"]),
    key=lambda value: len(value[0]), reverse=True,
)
_INLINE_LABELS = [
    (re.compile(r'(?<![A-Za-z])' + r'[\s.()/（）]*'.join(re.escape(c) for c in alias) + r'\s*[:：]', re.I), field_id)
    for alias, field_id, _ in _ALIAS_LOOKUP if len(alias) > 2 or not alias.isascii()
]


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _clean_cell(value: Any) -> str:
    """Normalize a spreadsheet cell while preserving line boundaries in addresses/items."""
    if value is None:
        return ""
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in str(value).splitlines()]
    return "\n".join(line for line in lines if line)


def _norm(value: str) -> str:
    return re.sub(_NORMALIZE, "", value.lower())


def _field_for(label: str) -> tuple[str, int] | None:
    label = re.sub(r"^\s*\d+\s*[.)、]\s*", "", label)
    compact = _norm(label)
    if not compact:
        return None
    for alias, field_id, _original in _ALIAS_LOOKUP:
        if alias and (compact == alias or compact.startswith(alias)):
            if len(alias) <= 3 and alias.isascii() and compact != alias and compact[len(alias)].isalpha() and compact[len(alias)].isascii():
                continue
            if field_id == "document_date" and len(compact) > len(alias) and alias == "date":
                continue
            return field_id, len(alias)
    return None


def _looks_like_field_label(value: str) -> bool:
    """识别被 OCR/合并单元格污染的字段标签，避免把它们作为数据值。"""
    if _clean(value).endswith((":", "：")):
        return True
    compact = _norm(value)
    if not compact:
        return False
    if compact in {'申报数', '申报总', '申报', '计量', '单位', 'lidtype', 'marks'}:
        return True
    # Bilingual Excel headers are sometimes read as a label prefix plus its
    # translated tail (for example "and Specifications 品名和规格"). Treat
    # common conjunctions followed by a field alias as label text too.
    tail = re.sub(r"^(?:and|or|및|和|与)+", "", compact)
    if tail != compact and _field_for(tail):
        return True
    for alias, _field_id, _original in _ALIAS_LOOKUP:
        position = compact.find(alias)
        if position == 0 and len(compact) - len(alias) <= 3:
            return True
    return _is_field_header(value) is not None


def _is_field_header(value: str) -> tuple[str, int] | None:
    """Match a complete header, including bilingual labels, but not label prefixes."""
    found = _field_for(value)
    if not found:
        return None
    field_id = found[0]
    value = re.sub(r'^\s*\d+\s*[.)、]\s*', '', value)
    remainder = _norm(value)
    aliases = [alias for alias, candidate_id, _ in _ALIAS_LOOKUP if candidate_id == field_id]
    for alias in aliases:
        remainder = remainder.replace(alias, "")
    remainder = re.sub(r'(?:usd|rmb|cny|eur|gbp|ctns?|cartons?|kgs?|千克|公斤|地区|20gp)', '', remainder)
    if field_id in {'unit_price', 'amount', 'quantity'}:
        remainder = re.sub(r'(?:fob|cif|cfr)[a-z]+', '', remainder)
    return found if not remainder else None


def _split_label_value(text: str) -> tuple[str, str] | None:
    text = re.sub(r'^\s*\d+\s*[.)、]\s*', '', text)
    lines = text.splitlines()
    if len(lines) > 1 and _is_field_header(lines[0]):
        field_id = _is_field_header(lines[0])[0]
        value = '\n'.join(lines[1:]).strip()
        if not _is_field_header(text) and _valid_field_value(field_id, value):
            return field_id, value
    # 先处理常见冒号标签；不依赖字段所在的固定行、列或坐标。
    for separator in (":", "：", "\t"):
        if separator in text:
            label, value = text.split(separator, 1)
            match = _is_field_header(label)
            if _norm(label) == 'loadingportdestinationport':
                match = ('loading_port', 0)
            if match and _clean(value):
                value = _clean_cell(value)
                # 表头常把多个双语字段挤在同一单元格，例如
                # "INVOICE NO. 发票号: INVOICE DATE 发票日期"。冒号后的内容
                # 仍然是字段标签，不能当作发票号候选。
                if not _looks_like_field_label(value):
                    return match[0], value
    if any(separator in text for separator in (":", "：", "\t")):
        # 有显式分隔符但标签/值未能可靠解析时，不再走无分隔符的前缀猜测。
        return None
    if _is_field_header(text):
        return None
    compact = _norm(text)
    for alias, field_id, _original in _ALIAS_LOOKUP:
        if alias and compact.startswith(alias) and len(text) > len(alias):
            if len(alias) <= 3 and alias.isascii() and compact != alias and compact[len(alias)].isascii() and compact[len(alias)].isalpha():
                continue
            # alias 已去掉空格和标点，不能直接用它的长度切原文；
            # 否则 "INVOICE NO." 会从 "INVOICE NO. 发票号" 中间切开，
            # 把残余标签误当成识别值。
            consumed = 0
            prefix = ""
            for consumed, character in enumerate(text, 1):
                prefix += _norm(character)
                if len(prefix) >= len(alias):
                    break
            match = _field_for(text[:consumed])
            if match and match[0] == field_id and _norm(text[:consumed]) == alias:
                value = text[consumed:].lstrip(" \t:：-—()（）")
                if _looks_like_field_label(value):
                    continue
                if value:
                    return field_id, _clean_cell(value)
    return None


def _rows_plain(text: str) -> list[TextRow]:
    return [TextRow(line.strip(), f"文本第 {index} 行") for index, line in enumerate(text.splitlines(), 1) if line.strip()]


def _read_xlsx(content: bytes) -> list[TextRow]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows: list[TextRow] = []
    for sheet in book.worksheets:
        for row_index, cells in enumerate(sheet.iter_rows(), 1):
            values = []
            for cell in cells:
                value = _clean_cell(cell.value)
                if value and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    value = _format_currency_cell(cell.value, cell.number_format)
                if value:
                    values.append((cell.coordinate, value))
            if values:
                rows.append(TextRow(" | ".join(value for _, value in values), f"工作表「{sheet.title}」第 {row_index} 行", values))
    book.close()
    return rows


def _read_xls(content: bytes) -> list[TextRow]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    rows: list[TextRow] = []
    for sheet in book.sheets():
        for row_index in range(sheet.nrows):
            values = [(f"{xlrd.colname(col_index)}{row_index + 1}", _clean_cell(sheet.cell_value(row_index, col_index)))
                      for col_index in range(sheet.ncols) if _clean_cell(sheet.cell_value(row_index, col_index))]
            if values:
                rows.append(TextRow(" | ".join(value for _, value in values), f"工作表「{sheet.name}」第 {row_index + 1} 行", values))
    return rows


def _read_docx(content: bytes) -> list[TextRow]:
    from docx import Document

    document = Document(io.BytesIO(content))
    rows = [TextRow(_clean(paragraph.text), f"段落 {index}")
            for index, paragraph in enumerate(document.paragraphs, 1) if _clean(paragraph.text)]
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            seen = set()
            values = []
            for col_index, cell in enumerate(row.cells, 1):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                if _clean_cell(cell.text):
                    values.append((f"表格 {table_index} 第 {row_index} 行第 {col_index} 列", _clean_cell(cell.text)))
            if values:
                rows.append(TextRow(" | ".join(value for _, value in values), f"表格 {table_index} 第 {row_index} 行", values))
    return rows


def _read_doc(content: bytes) -> list[TextRow]:
    executable = os.environ.get('LIBREOFFICE_PATH') or shutil.which('soffice')
    if not executable:
        for base in [os.environ.get('ProgramFiles', ''), os.environ.get('ProgramFiles(x86)', '')]:
            candidate = Path(base) / 'LibreOffice/program/soffice.exe'
            if candidate.is_file():
                executable = str(candidate)
                break
    if not executable:
        from server.document_sources import render_word
        return _read_docx(render_word(content, '.doc', 'docx'))
    with tempfile.TemporaryDirectory(prefix='trade-doc-') as directory:
        source = Path(directory) / 'input.doc'
        source.write_bytes(content)
        subprocess.run([executable, '-env:UserInstallation=' + (Path(directory) / 'profile').as_uri(),
                        '--headless', '--convert-to', 'docx', '--outdir', directory, str(source)],
                       capture_output=True, timeout=90,
                       creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        converted = source.with_suffix('.docx')
        if not converted.is_file():
            raise ValueError('旧版 Word 转换失败，请确认文件没有损坏或加密。')
        return _read_docx(converted.read_bytes())


def _ocr_engine():
    try:
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise RuntimeError("本机 OCR 组件未安装；请安装 rapidocr 与 onnxruntime 后重启本地服务。") from exc
    if not hasattr(_ocr_engine, "instance"):
        _ocr_engine.instance = RapidOCR(params={
            'EngineConfig.onnxruntime.intra_op_num_threads': 2,
            'EngineConfig.onnxruntime.inter_op_num_threads': 1,
        })
    return _ocr_engine.instance


def _ocr_image(image: Any, location: str, native_words: list | None = None) -> list[TextRow]:
    """Keep native OCR failures and Windows crash dialogs out of the web host."""
    buffer = io.BytesIO()
    image.convert('RGB').save(buffer, format='PNG')
    payload = json.dumps({'image': base64.b64encode(buffer.getvalue()).decode('ascii'),
                          'location': location, 'native_words': native_words}, ensure_ascii=False).encode('utf-8')
    try:
        result = subprocess.run([sys.executable, str(Path(__file__).with_name('trade_ocr_worker.py'))],
                                input=payload, capture_output=True, timeout=120,
                                creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError('本机 OCR 超时，请减少扫描页数或降低图片分辨率后重试。') from exc
    if result.returncode:
        raise RuntimeError(f'本机 OCR 运行失败（退出码 {result.returncode}）。请检查 OCR 依赖；Windows 已验证使用 onnxruntime 1.20.1。')
    try:
        data = json.loads(result.stdout.decode('utf-8').splitlines()[-1])
        return [TextRow(**row) for row in data]
    except (ValueError, TypeError, IndexError) as exc:
        raise RuntimeError('本机 OCR 返回内容无效，请检查 OCR 组件。') from exc


def _ocr_image_local(image: Any, location: str, native_words: list | None = None) -> list[TextRow]:
    import numpy as np
    result = _ocr_engine()(np.asarray(image.convert("RGB")))
    items = getattr(result, "txts", None)
    scores = getattr(result, "scores", None)
    boxes = getattr(result, "boxes", None)
    if items is None and result:
        # 兼容 RapidOCR 旧版返回 list[(box, text, score)] 的结构。
        items = [entry[1] for entry in result]
        scores = [entry[2] for entry in result]
        boxes = [entry[0] for entry in result]
    rows: list[TextRow] = []
    positioned = []
    for index, text in enumerate(items if items is not None else []):
        score = float(scores[index]) if scores is not None and index < len(scores) else None
        box = boxes[index] if boxes is not None and index < len(boxes) else None
        if native_words and box is not None:
            left, top = float(min(p[0] for p in box)), float(min(p[1] for p in box))
            right, bottom = float(max(p[0] for p in box)), float(max(p[1] for p in box))
            native = [w for w in native_words if left - 4 <= w[0] + w[2] <= right + 4 and top - 4 <= w[1] + w[3] <= bottom + 4]
            if native:
                exact = ' '.join(w[4] for w in sorted(native, key=lambda w: (round(w[1] / 4), w[0])))
                prefix = re.match(r'^(.+?[:：])', text)
                if prefix and ':' not in exact and '：' not in exact:
                    exact = prefix.group(1) + exact
                text = exact
        position = ""
        if box is not None:
            try:
                position = "，区域 " + ",".join(str(round(float(v))) for point in box for v in point)
            except (TypeError, ValueError):
                pass
        text = _clean(text)
        if text:
            if box is not None:
                positioned.append((float(box[0][0]), float(box[0][1]), float(box[2][1] - box[0][1]), text, location + position, score))
            else:
                rows.append(TextRow(text, location + position, ocr_score=score))
    lines: list[list[Any]] = []
    for entry in sorted(positioned, key=lambda v: (v[1], v[0])):
        line = next((line for line in reversed(lines[-3:]) if abs(line[0][1] - entry[1]) < max(5, min(line[0][2], entry[2]) * 0.5)), None)
        if line is None:
            lines.append([entry])
        else:
            line.append(entry)
    for index, line in enumerate(lines, 1):
        line.sort(key=lambda v: v[0])
        cells = [(f'{entry[4]}第 {i} 列', entry[3]) for i, entry in enumerate(line, 1)]
        score_values = [entry[5] for entry in line if entry[5] is not None]
        rows.append(TextRow(' | '.join(v for _, v in cells), f'{location}第 {index} 行', cells,
                            min(score_values) if score_values else None,
                            {loc: entry[0] for (loc, _), entry in zip(cells, line)}))
    return rows


def _read_pdf(content: bytes) -> list[TextRow]:
    import pymupdf as fitz

    document = fitz.open(stream=content, filetype="pdf")
    rows: list[TextRow] = []
    pending_header = None
    for page_index, page in enumerate(document, 1):
        start = len(rows)
        page_text = page.get_text("text").strip()
        # Broken ToUnicode maps can produce nonempty but unreadable text.
        unusual = sum(1 for c in page_text if '\u0400' <= c <= '\u1fff' or c == '\ufffd')
        if page_text and unusual / max(1, len(page_text)) < 0.08:
            words = page.get_text('words')
            tables = page.find_tables().tables
            text_words = [w for w in words if not any(
                t.bbox[0] <= (w[0] + w[2]) / 2 <= t.bbox[2] and
                t.bbox[1] <= (w[1] + w[3]) / 2 <= t.bbox[3] for t in tables)]
            lines: list[list[Any]] = []
            for word in sorted(text_words, key=lambda w: (w[1], w[0])):
                line = next((line for line in reversed(lines[-3:]) if abs(line[0][1] - word[1]) < 4), None)
                if line is None:
                    lines.append([word])
                else:
                    line.append(word)
            for line_index, line in enumerate(lines, 1):
                segments: list[str] = []
                right = None
                for word in sorted(line, key=lambda w: w[0]):
                    if right is None or word[0] - right > 12:
                        segments.append(word[4])
                    else:
                        segments[-1] += ' ' + word[4]
                    right = word[2]
                location = f'第 {page_index} 页第 {line_index} 行'
                cells = [(f'{location}第 {i} 列', text) for i, text in enumerate(segments, 1)]
                rows.append(TextRow(' | '.join(segments), location, cells))
            for table_index, table in enumerate(tables, 1):
                table_rows = table.extract()
                if pending_header and table_index == 1 and table_rows and len(table_rows[0]) == len(pending_header):
                    first = table_rows[0]
                    if sum(_is_field_header(_clean_cell(v)) is not None for v in first) < 2:
                        table_rows[0] = [''.join((a or '').split()) + ''.join((b or '').split()) for a, b in zip(pending_header, first)]
                pending_header = None
                if table_rows and sum(_field_for(_clean_cell(v)) is not None for v in table_rows[-1]) >= 3:
                    pending_header = table_rows[-1]
                for row_index, cells in enumerate(table_rows, 1):
                    location = f'第 {page_index} 页表格 {table_index} 第 {row_index} 行'
                    values = [(f'{location}第 {i} 列', _clean_cell(value)) for i, value in enumerate(cells, 1) if value]
                    if values:
                        # Wrapped numbers in PDF tables are a single cell value.
                        values = [(loc, re.sub(r'\s+', '', v) if re.fullmatch(r'[\d.\s]+', v) else v) for loc, v in values]
                        rows.append(TextRow(' | '.join(v for _, v in values), location, values))
            # Some PDF forms paint labels as vector outlines. Text extraction
            # then returns only filled-in values, so OCR the visible form too.
            if tables and len(_extract_candidates(rows[start:])) < 3:
                from PIL import Image
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                rows.extend(_ocr_image(Image.open(io.BytesIO(pixmap.tobytes('png'))), f'第 {page_index} 页 OCR', words))
        else:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            from PIL import Image
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            rows.extend(_ocr_image(image, f"第 {page_index} 页 OCR"))
    document.close()
    return rows


def _read_image(content: bytes) -> list[TextRow]:
    from PIL import Image
    return _ocr_image(Image.open(io.BytesIO(content)), "图片 OCR")


def _read_document(filename: str, content: bytes) -> tuple[list[TextRow], str]:
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        return _read_xlsx(content), "文本/表格解析"
    if extension == ".xls":
        return _read_xls(content), "文本/表格解析"
    if extension == ".docx":
        return _read_docx(content), "文本/表格解析"
    if extension == ".pdf":
        rows = _read_pdf(content)
        method = "文字层解析" if not any("OCR" in row.location for row in rows) else "文字层解析 + 本地 OCR"
        return rows, method
    if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        return _read_image(content), "本地 OCR"
    if extension in {".txt", ".csv"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("gb18030", errors="replace")
        return _rows_plain(text), "文本解析"
    if extension == ".doc":
        return _read_doc(content), '旧版 Word 本地转换/表格解析'
    raise ValueError("支持的格式：PDF、DOCX、XLSX、XLS、TXT、CSV、PNG、JPG、TIFF。")


def _match_table_headers(rows: list[TextRow]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows[:-1]):
        if not row.cells:
            continue
        # Inline forms store label and value in the same row (label cells end
        # with a colon); they must not be treated as column headers for the next row.
        if any(re.search(r"[:：]\s*$", cell_text) for _location, cell_text in row.cells):
            continue
        mapped: list[tuple[str, str, str]] = []
        for cell_location, cell_text in row.cells:
            found = _is_field_header(cell_text)
            if found:
                mapped.append((found[0], cell_location, cell_text))
        if len(mapped) < 2 or not _is_column_header_row(row):
            continue
        # 表头与数据按列关联；跳过同列的单位/注释，数值字段继续向下找数值。
        for field_id, location, header in mapped:
            value = ""
            value_location = location
            column = _column_key(location)
            seen_data = False
            for next_row in rows[row_index + 1:]:
                if _row_scope(next_row.location) != _row_scope(row.location):
                    break
                if _is_column_header_row(next_row) or re.match(r'^(?:TOTAL|合计|总计)', next_row.text, re.I):
                    break
                next_values = {_column_key(cell_location): cell_value for cell_location, cell_value in next_row.cells}
                positions = row.cell_x
                value_locations = {_column_key(loc): loc for loc, _ in next_row.cells}
                if positions and next_row.cell_x:
                    next_values = {}
                    value_locations = {}
                    for loc, val in next_row.cells:
                        nearest = min(positions, key=lambda h: abs(positions[h] - next_row.cell_x[loc]))
                        key = _column_key(nearest)
                        next_values[key] = val
                        value_locations[key] = loc
                weight_continuation = _norm(header) == '数重量' and re.fullmatch(r'\d+(?:\.\d+)?\s*(?:KG|KGS|千克)', next_values.get(column, ''), re.I)
                if weight_continuation:
                    candidates.append({'fieldId': 'declared_weight', 'value': next_values[column], 'location': value_locations[column], 'rawText': f'{header}: {next_values[column]}', 'confidence': 0.8})
                    continue
                package_continuation = field_id == 'package_type' and re.fullmatch(r'/箱/(\d+)', next_values.get(column, ''))
                if package_continuation:
                    candidates.append({'fieldId': 'package_count', 'value': package_continuation.group(1), 'location': value_locations[column], 'rawText': f'{header}: {next_values[column]}', 'confidence': 0.8})
                    continue
                valid_count = sum(_valid_field_value(fid, next_values.get(_column_key(loc), '')) for fid, loc, _ in mapped)
                if 'OCR' in row.location and valid_count < 2:
                    if seen_data:
                        break
                    continue
                if valid_count >= 2:
                    seen_data = True
                candidate_value = next_values.get(column, "")
                if _valid_field_value(field_id, candidate_value):
                    value = candidate_value
                    value_location = value_locations[column]
                    sheet = _sheet_name(row.location)
                    evidence_location = f"工作表「{sheet}」!{value_location}" if sheet else value_location
                    candidates.append({"fieldId": field_id, "value": value, "location": evidence_location, "rawText": f"{header}: {value}", "confidence": 0.82})
                    continue
                if candidate_value and _looks_like_field_label(candidate_value):
                    break
    return candidates


def _is_column_header_row(row: TextRow) -> bool:
    if any(value.rstrip().endswith((':', '：')) for _, value in row.cells):
        return False
    headers = sum(_is_field_header(value) is not None for _, value in row.cells)
    numeric = any(re.match(r'^[$¥￥€£]?\s*\d', value) for _, value in row.cells)
    return headers >= 2 and not numeric and headers / max(1, len(row.cells)) >= 0.5


def _row_scope(location: str) -> str:
    return _sheet_name(location) or re.split(r'第\s*\d+\s*行', location)[0]


def _column_key(location: str) -> str:
    cell = re.match(r"([A-Z]+)\d+$", location)
    if cell:
        return cell.group(1)
    column = re.search(r"第\s*(\d+)\s*列$", location)
    return column.group(1) if column else location


def _sheet_name(location: str) -> str:
    match = re.search(r"工作表「([^」]+)」", location)
    return match.group(1) if match else ""


def _is_unit_text(value: str) -> bool:
    return bool(re.fullmatch(r"\s*(?:/\s*)?(?:CTNS?|CARTONS?|PCS?|KGS?|KG|G|GR|CBM|SETS?|TONS?)\s*", value, re.I))


def _format_currency_cell(value: int | float, number_format: str) -> str:
    """Keep an explicit currency symbol from Excel's number format in candidates."""
    first_section = number_format.split(";", 1)[0]
    if '%' in first_section:
        fraction = re.search(r'\.([0#]+)', first_section)
        return f'{value * 100:.{len(fraction.group(1)) if fraction else 0}f}%'
    symbol_match = re.search(r"\\?([$¥￥€£])|\[\$([A-Z]{3})", first_section)
    if not symbol_match:
        return str(value)
    symbol = symbol_match.group(1) or symbol_match.group(2)
    decimals = 0
    decimal_part = first_section.split(".", 1)
    if len(decimal_part) == 2:
        decimals = len(re.match(r"[0#]+", decimal_part[1]).group(0)) if re.match(r"[0#]+", decimal_part[1]) else 0
    number = f"{value:,.{decimals}f}"
    return f"{symbol}{number}"


def _valid_field_value(field_id: str, value: str) -> bool:
    if not value or _looks_like_field_label(value) or _is_unit_text(value):
        return False
    if field_id in {"quantity", "quantity_per_container", "sample_quantity", "declared_weight", "package_count", "amount", "unit_price", "gross_weight", "net_weight"}:
        return bool(re.fullmatch(r"\s*(?:[$¥￥€£]\s*|(?:USD|RMB|CNY|EUR|GBP)\s*)?[\d,]+(?:\.\d+)?\s*(?:KGS?|千克|公斤|CTNS?|CARTONS?|箱|件|罐|个|吨|美元|美金|USD|EUR|CNY)?\s*(?:[×x*]\s*\d+\s*(?:罐|个))?\s*", value, re.I))
    if field_id == 'currency':
        return bool(re.fullmatch(r'USD|CNY|RMB|EUR|GBP|JPY|美元|美金|人民币|欧元|日元|英镑|[$¥￥€£]', value, re.I))
    if field_id == 'hs_code':
        return bool(re.fullmatch(r'\d{6,10}(?:\.\d+)*(?:\([^\n]*)?', value))
    if field_id == 'container_count':
        return bool(re.search(r'\d', value))
    if field_id in {'document_date', 'shipment_date', 'production_date', 'cutoff_date'}:
        return bool(re.search(r'\d', value))
    return True


def _split_extracted_value(field_id: str, value: str) -> list[tuple[str, str]]:
    if field_id in {'amount', 'unit_price'}:
        currency = re.search(r'USD|EUR|CNY|RMB|GBP|美元|美金|人民币|欧元|[$¥￥€£]', value, re.I)
        if currency:
            return [(field_id, value), ('currency', currency.group(0))]
    if field_id == 'hs_code':
        return [(field_id, value.split('(', 1)[0])]
    if field_id in {"seller_name", "buyer_name"}:
        lines = [line.strip() for line in value.splitlines() if line.strip()]
        if len(lines) > 1:
            name = lines[0]
            address = "\n".join(lines[1:])
            address = re.sub(r"^ADD\s*:\s*", "", address, flags=re.I)
            address_id = "seller_address" if field_id == "seller_name" else "buyer_address"
            return [(field_id, name), (address_id, address)]
        add_match = re.search(r"\s+ADD\s*:\s*", value, re.I)
        if add_match and field_id == "seller_name":
            return [(field_id, value[:add_match.start()].strip()), ("seller_address", value[add_match.end():].strip())]
    if field_id == "item_description":
        code = re.search(r"\bH\s*\.?\s*S\s*\.?\s*CODE\s*[:：]?\s*([0-9][0-9.\-/]*)", value, re.I)
        specification = re.search(r"\b\d+\s*[xX×]\s*\d+\s*/\s*\d+\s*(?:g|gr|kg)?\b", value, re.I)
        result = [(field_id, value)]
        if specification:
            result.append(("item_specification", specification.group(0).replace(" ", "")))
        if code:
            result.append(("hs_code", code.group(1)))
        if specification or code:
            description_lines = []
            for line in value.splitlines():
                if specification:
                    line = re.sub(re.escape(specification.group(0)), '', line, flags=re.I)
                if code:
                    line = line.replace(code.group(0), '')
                description_lines.append(line.strip())
            result[0] = (field_id, _clean(" ".join(line for line in description_lines if line)))
        return result
    if field_id == "loading_port":
        match = re.match(r"\s*(.+?)\s+to\s+(.+?)(?:\s+by\s+sea)?\s*$", value, re.I)
        if match:
            return [("loading_port", match.group(1).strip()), ("destination_port", match.group(2).strip())]
    return [(field_id, value)]


def _multiple_labels(text: str) -> list[tuple[str, str]]:
    matches = []
    for pattern, field_id in _INLINE_LABELS:
        matches.extend((m.start(), m.end(), field_id) for m in pattern.finditer(text))
    selected = []
    for start, end, field_id in sorted(matches, key=lambda m: (m[0], -m[1])):
        if not selected or start >= selected[-1][1]:
            selected.append((start, end, field_id))
    if len(selected) < 2:
        return []
    return [(field_id, text[end:selected[i + 1][0] if i + 1 < len(selected) else len(text)].strip())
            for i, (_, end, field_id) in enumerate(selected)]


def _customs_detail(text: str) -> list[tuple[str, str]]:
    match = re.match(r'^\s*(\d{8,10})\s+(.+?)\s+(\d+(?:\.\d+)?\s*KG)\s+(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)\s*/\s*(美金|美元|USD|CNY|人民币)', text, re.I)
    if not match:
        return []
    return list(zip(['hs_code', 'item_description', 'net_weight', 'unit_price', 'amount', 'currency'], match.groups()))


def _candidate_key(field_id: str, value: str) -> str:
    if field_id in {'quantity', 'quantity_per_container', 'sample_quantity', 'package_count', 'unit_price', 'amount', 'net_weight', 'gross_weight', 'declared_weight'}:
        text = re.sub(r'\s+|,', '', value.lower())
        text = re.sub(r'\d+(?:\.\d+)?', lambda m: format(Decimal(m.group(0)).normalize(), 'f'), text)
        text = re.sub(r'公斤|千克|kgs', 'kg', text)
        return re.sub(r'美金|美元', 'usd', text)
    return _norm(value)


def _extract_candidates(rows: list[TextRow]) -> dict[str, list[dict[str, Any]]]:
    found: dict[str, list[dict[str, Any]]] = {}
    for row_index, row in enumerate(rows):
        for loc, text in row.cells:
            header = _is_field_header(text.rstrip(':：'))
            if header and header[0] in {'unit_price', 'amount'}:
                for fid, pattern in [('currency', r'\b(?:USD|EUR|CNY|RMB|GBP)\b'), ('trade_term', r'\b(?:FOB|CIF|CFR|EXW|FCA|DAP|DDP)\b')]:
                    match = re.search(pattern, text, re.I)
                    if match:
                        found.setdefault(fid, []).append({'fieldId': fid, 'value': match.group(0), 'location': row.location, 'rawText': text, 'confidence': 0.94})
            if header and header[0] == 'composition':
                parts = []
                for ingredient in rows[row_index + 1:row_index + 31]:
                    cells = ingredient.cells
                    if _row_scope(ingredient.location) != _row_scope(row.location) or len(cells) != 2 or not re.fullmatch(r'\d+(?:\.\d+)?%?', cells[1][1]):
                        break
                    parts.append(re.sub(r'^\s*\d+\s*[.)、]\s*', '', cells[0][1]) + ': ' + cells[1][1])
                if parts:
                    found.setdefault('composition', []).append({'fieldId': 'composition', 'value': '\n'.join(parts), 'location': row.location, 'rawText': '\n'.join(parts), 'confidence': 0.94})
        if _is_column_header_row(row):
            continue
        for cell_index, (cell_location, cell_text) in enumerate(row.cells or [(row.location, row.text)]):
            parsed_values = _multiple_labels(cell_text) or _customs_detail(cell_text)
            parsed = parsed_values[0] if parsed_values else _split_label_value(cell_text)
            if parsed is None:
                header_text = cell_text.rstrip(" \t:：")
                alias = _is_field_header(header_text)
                if alias and row.cells and cell_index + 1 < len(row.cells):
                    next_cell = row.cells[cell_index + 1][1]
                    if _valid_field_value(alias[0], next_cell):
                        parsed = (alias[0], next_cell)
            if parsed is None:
                continue
            field_id, value = parsed
            if not _valid_field_value(field_id, value):
                continue
            # 日期、金额等字段误捕获整行正文时不作候选。
            if len(value) > 600:
                continue
            score = row.ocr_score
            confidence = max(0.4, min(0.99, score * 0.85 + 0.1)) if score is not None else 0.94
            expanded = [pair for fid, val in (parsed_values or [(field_id, value)]) for pair in _split_extracted_value(fid, val)]
            for extracted_field, extracted_value in expanded:
                if not _valid_field_value(extracted_field, extracted_value):
                    continue
                sheet = _sheet_name(row.location)
                evidence_location = f"工作表「{sheet}」!{cell_location}" if sheet and re.fullmatch(r"[A-Z]{1,3}[1-9]\d*", cell_location) else cell_location
                item = {"fieldId": extracted_field, "value": extracted_value, "location": evidence_location, "rawText": cell_text, "confidence": round(confidence, 2)}
                found.setdefault(extracted_field, []).append(item)
    for item in _match_table_headers(rows):
        for extracted_field, extracted_value in _split_extracted_value(item["fieldId"], item["value"]):
            if _valid_field_value(extracted_field, extracted_value):
                found.setdefault(extracted_field, []).append({**item, "fieldId": extracted_field, "value": extracted_value})
    # 同一文件中重复出现的相同值保留来源但合并候选；不同值作为冲突候选显示。
    for field_id, items in list(found.items()):
        unique: dict[str, dict[str, Any]] = {}
        for item in items:
            key = _candidate_key(field_id, item["value"])
            if not key:
                continue
            if key not in unique:
                unique[key] = {**item, 'normalizedValue': key, "evidence": []}
                if field_id.endswith('_date') and re.search(r'\d{4}[-/]\d{1,2}[-/.]{2,}', item['value']):
                    unique[key]['needsReview'] = True
                    unique[key]['reviewNote'] = '原文日期格式异常，请核对源文件后确认。'
            unique[key]["evidence"].append({"location": item["location"], "rawText": item["rawText"]})
            unique[key]["confidence"] = max(unique[key]["confidence"], item["confidence"])
        found[field_id] = list(unique.values())
    return found


@router.get("/fields")
async def document_fields():
    return {"fields": FIELD_DEFS}


@router.post("/fields/export")
async def export_document_fields(payload: dict[str, Any]):
    """生成按网页字段分类的可编辑工作簿。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    values = payload.get("values") or {}
    custom_fields = payload.get("customFields") or []
    products = payload.get("products") or payload.get("goods") or []
    definitions = {field["id"]: field for field in FIELD_DEFS}
    book = Workbook()
    book.remove(book.active)
    groups = list(dict.fromkeys([field["group"] for field in FIELD_DEFS] + ["自定义"]))
    for group in groups:
        sheet = book.create_sheet(group[:31])
        is_product = group == "产品"
        headers = (["产品序号"] if is_product else []) + ["字段标识", "字段名称", "字段值"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3151CB")
        if is_product:
            for index, item in enumerate(products, 1):
                item_values = item.get("values", item)
                for field in FIELD_DEFS:
                    if field["group"] == group and field["id"] in REUSABLE_PRODUCT_FIELDS:
                        sheet.append([index, field["id"], field["label"], item_values.get(field["id"], "")])
        elif group == "自定义":
            for index, field in enumerate(custom_fields, 1):
                sheet.append([f"custom-{index}", field.get("label", ""), field.get("value", "")])
        else:
            for field in FIELD_DEFS:
                if field["group"] == group:
                    sheet.append([field["id"], field["label"], values.get(field["id"], "")])
        for col, width in enumerate(([14] if is_product else []) + [24, 28, 48], 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    book.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''trade-document-fields.xlsx"})


@router.post("/fields/import")
async def import_document_fields(file: UploadFile = File(...)):
    """按字段标识/字段名读取工作簿，兼容列错位及传统三列表格。"""
    from openpyxl import load_workbook

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "字段工作簿不能超过 10 MB。")
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"无法读取 Excel 工作簿：{exc}") from exc
    by_id = {field["id"].lower(): field for field in FIELD_DEFS}
    by_label = {re.sub(_NORMALIZE, "", field["label"].lower()): field for field in FIELD_DEFS}
    values: dict[str, str] = {}
    products_by_index: dict[str, dict[str, str]] = {}
    automatic_item = 1
    custom_fields: list[dict[str, str]] = []
    row_count = 0
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            header_norm = [re.sub(_NORMALIZE, "", str(value or "").lower()) for value in header]
            id_col = next((i for i, value in enumerate(header_norm) if value in {"字段标识", "字段id", "fieldid", "id"}), None)
            label_col = next((i for i, value in enumerate(header_norm) if value in {"字段名称", "字段名", "字段", "fieldname", "label"}), None)
            value_col = next((i for i, value in enumerate(header_norm) if value in {"字段值", "值", "内容", "value"}), None)
            item_col = next((i for i, value in enumerate(header_norm) if value in {"产品序号", "货物序号", "商品序号", "itemno", "itemindex"}), None)
            for row in rows:
                cells = ["" if value is None else str(value).strip() for value in row]
                if not any(cells):
                    continue
                row_count += 1
                # Use named columns where available. Otherwise classify cells by content,
                # allowing the ID, label, and value to move to any column.
                field = None
                id_value = cells[id_col].lower() if id_col is not None and id_col < len(cells) else ""
                label_value = cells[label_col] if label_col is not None and label_col < len(cells) else ""
                field = by_id.get(id_value) or by_label.get(re.sub(_NORMALIZE, "", label_value.lower()))
                if field is None:
                    for cell in cells:
                        field = by_id.get(cell.lower()) or by_label.get(re.sub(_NORMALIZE, "", cell.lower()))
                        if field:
                            break
                if sheet.title == "自定义" or (id_value.startswith("custom-") and label_value):
                    if label_value:
                        custom_fields.append({"label": label_value, "value": cells[value_col] if value_col is not None and value_col < len(cells) else ""})
                    continue
                if not field:
                    continue
                if value_col is not None and value_col < len(cells) and cells[value_col]:
                    field_value = cells[value_col]
                else:
                    recognized = {field["id"].lower(), re.sub(_NORMALIZE, "", field["label"].lower())}
                    remaining = [cell for cell in cells if cell and cell.lower() not in recognized]
                    item_value = cells[item_col] if item_col is not None and item_col < len(cells) else ""
                    remaining = [cell for cell in remaining if cell != item_value]
                    field_value = remaining[-1] if remaining else ""
                if field["group"] == "产品":
                    if item_col is not None and item_col < len(cells) and cells[item_col]:
                        item_index = cells[item_col]
                    else:
                        item_index = str(automatic_item)
                        if field["id"] in products_by_index.setdefault(item_index, {}):
                            automatic_item += 1
                            item_index = str(automatic_item)
                    products_by_index.setdefault(item_index, {})[field["id"]] = field_value
                else:
                    values[field["id"]] = field_value
    finally:
        book.close()
    products = [products_by_index[key] for key in sorted(products_by_index, key=lambda value: (not value.isdigit(), int(value) if value.isdigit() else value))]
    if not values and not products and not custom_fields:
        raise HTTPException(400, "没有识别到字段。请确保工作表中包含字段标识或可识别的字段名称列。")
    return {"values": values, "products": products, "goods": products, "customFields": custom_fields, "rowsRead": row_count}


@router.post("/extract")
async def extract_document(
    file: UploadFile = File(...),
):
    filename = file.filename or "uploaded-file"
    content = await file.read()
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(413, "单个文件不能超过 25 MB。")
    try:
        rows, method = _read_document(filename, content)
        candidates = _extract_candidates(rows)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(503, str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 - 向用户给出文件级解析错误
        raise HTTPException(400, f"无法读取该文件：{exc}") from exc
    return {
        "document": {"fileName": filename, "method": method, "textRows": len(rows)},
        "fields": candidates,
        'warnings': [{'fieldId': field_id, 'message': item['reviewNote'], 'location': item['location']}
                     for field_id, items in candidates.items() for item in items if item.get('needsReview')],
        "unmatchedText": [
            {"text": row.text[:500], "location": row.location}
            for row in rows
            if not any(row.text in candidate.get("rawText", "") for values in candidates.values() for candidate in values)
        ][:80],
    }
