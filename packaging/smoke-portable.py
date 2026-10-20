"""Exercise the actual EXE's HTTP and multiprocessing paths after a build."""
import json
import hashlib
import io
import sys
import time
from pathlib import Path

import httpx


def check_documents(client):
    from openpyxl import Workbook
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    import pymupdf

    text = 'Invoice No: TEST20261020\nQuantity: 75'
    book = Workbook()
    book.active.append(['Invoice No', 'TEST20261020'])
    book.active.append(['Quantity', 75])
    excel = io.BytesIO()
    book.save(excel)
    word = Document()
    for line in text.splitlines():
        word.add_paragraph(line)
    docx = io.BytesIO()
    word.save(docx)
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), text)
    pdf_bytes = pdf.tobytes()
    image = Image.new('RGB', (1400, 400), 'white')
    ImageDraw.Draw(image).multiline_text((50, 50), text, fill='black',
        font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 50), spacing=30)
    png = io.BytesIO()
    image.save(png, format='PNG')
    for name, content in [('invoice.xlsx', excel.getvalue()), ('invoice.docx', docx.getvalue()),
                          ('invoice.pdf', pdf_bytes), ('invoice.png', png.getvalue())]:
        response = client.post('/api/documents/extract', files={'file': (name, content)}, timeout=150)
        assert response.status_code == 200, (name, response.text[:500])
        fields = response.json()['fields']
        assert any(str(candidate['value']) == '75' for candidate in fields.get('quantity', [])), (name, fields)
        assert any(candidate.get('location') and candidate.get('rawText')
                   for candidate in fields['quantity']), name
    response = client.post('/api/documents/workbook-sessions', files={'file': ('invoice.xlsx', excel.getvalue())})
    assert response.status_code == 200, response.text
    base = '/api/documents/workbook-sessions/' + response.json()['sessionId']
    assert client.get(base).json()['sheets'][0]['rows'][1][1] == 75
    assert client.get(base + '/original').content == excel.getvalue()
    response = client.post('/api/documents/source-sessions', files={'file': ('invoice.pdf', pdf_bytes)})
    assert response.status_code == 200, response.text
    base = '/api/documents/source-sessions/' + response.json()['sessionId']
    focused = client.post(base + '/focus', json={'location': '第 1 页', 'text': 'TEST20261020'})
    assert focused.json()['matches'], focused.text
    assert client.get(base + '/pages/1.png').content.startswith(b'\x89PNG')
    assert client.get(base + '/original').content == pdf_bytes
    # Every static resource must be exactly the version being released.
    web_root = Path(__file__).resolve().parents[1] / 'application' / 'web'
    files = list(web_root.rglob('*'))
    count = 0
    for source in files:
        if source.is_file():
            response = client.get('/' + source.relative_to(web_root).as_posix())
            assert response.status_code == 200, source.name
            assert hashlib.sha256(response.content).digest() == hashlib.sha256(source.read_bytes()).digest(), source.name
            count += 1
    return dict(document_formats=['xlsx', 'docx', 'pdf', 'png OCR'], source_roundtrip=True, static_files=count)


def main():
    port = (Path(sys.argv[1]) / 'data' / 'port.txt').read_text().strip()
    with httpx.Client(base_url=f'http://127.0.0.1:{port}', timeout=60) as client:
        assert client.get('/api/health').json() == {'status': 'ok'}
        page = client.get('/')
        assert page.status_code == 200 and '最大装填' in page.text
        rows = [dict(l=.49,w=.4,h=.09,n=44,wt=18,type='纸箱',desc='法兰盘配件'),
                dict(l=.2,w=.2,h=.4,n=0,wt=1472,type='木箱',desc='主机箱体')]
        body = dict(container=[5.8,2.35,2.35], rows=rows, mode='standard',
                    budget=10, bottom_metric='none', threshold=100)
        snap = client.post('/api/jobs', json=body)
        assert snap.status_code == 202, snap.text
        job_id = snap.json()['id']
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            snap = client.get(f'/api/jobs/{job_id}').json()
            if snap['status'] not in ('queued','running'):
                break
            time.sleep(.25)
        assert snap['status'] == 'done', snap
        result = snap['result']
        assert result['fits'] and result['maxFlags'] == [False, True]
        assert result['stats']['placed'] >= 1798, result['stats']
        assert result['plan']['boxes'][0][3] == 44
        export = client.post('/api/plan/xlsx', json=dict(plan=result['plan'],stats=result['stats']))
        assert export.status_code == 200, export.text
        restored = client.post('/api/plan/import', files={'file': ('plan.xlsx',export.content)})
        assert restored.status_code == 200, restored.text
        assert restored.json()['result']['stats']['placed'] == result['stats']['placed']
        # The n=0 input must also survive saving and restoring an unfinished list.
        saved = client.post('/api/cargo/export?fmt=xlsx', json=body)
        restored_input = client.post('/api/cargo/import', files={'file': ('cargo.xlsx',saved.content)})
        assert [r['n'] for r in restored_input.json()['rows']] == [44, 0]
        documents = check_documents(client)
        print(json.dumps(dict(port=port, stats=result['stats'], quantities=[b[3] for b in result['plan']['boxes']],
                              xlsx_roundtrip=True, blank_quantity_roundtrip=True, documents=documents), ensure_ascii=False))


if __name__ == '__main__':
    main()
