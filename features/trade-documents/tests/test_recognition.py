"""Synthetic regressions; real business files are read only by the audit command."""
import io
from pathlib import Path
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / 'application'))
from server.trade_documents import TextRow, _read_document, _extract_candidates


def values(rows):
    return {key: [v['value'] for v in candidates]
            for key, candidates in _extract_candidates(rows).items()}


class RecognitionTests(unittest.TestCase):
    def test_malformed_source_date_is_preserved_and_flagged(self):
        candidate = _extract_candidates([TextRow('出口日期\n2026-10.--21', 'A1')])['shipment_date'][0]
        self.assertEqual(candidate['value'], '2026-10.--21')
        self.assertTrue(candidate.get('needsReview'))

    def test_decimal_points_do_not_merge_different_amounts(self):
        found = values([TextRow('Amount: 1.2', 'A1'), TextRow('Amount: 12', 'A2'), TextRow('Amount: 1.20', 'A3')])
        self.assertEqual(found.get('amount'), ['1.2', '12'])

    def test_declaration_composition_keeps_percent_display(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(['成分含量：'])
        sheet.append(['1)TEST CORN', .9])
        sheet['B2'].number_format = '0.00%'
        output = io.BytesIO()
        book.save(output)
        rows, _ = _read_document('test.xlsx', output.getvalue())
        self.assertEqual(values(rows).get('composition'), ['TEST CORN: 90.00%'])

    def test_native_ocr_crash_is_reported_without_crashing_host(self):
        from unittest.mock import patch
        from types import SimpleNamespace
        from PIL import Image
        import server.trade_documents as module
        with patch.object(module.subprocess, 'run', return_value=SimpleNamespace(returncode=3221225477, stdout=b'', stderr=b'')):
            with self.assertRaisesRegex(RuntimeError, 'OCR'):
                module._ocr_image(Image.new('RGB', (30, 30)), 'test')

    def test_inspection_weight_and_amount_with_units(self):
        rows = [TextRow('', '第 1 页 OCR第 1 行', [('第 1 列', '货物名称(中/外文)'), ('第 2 列', 'H.S.编码'), ('第 3 列', '产地'), ('第 4 列', '数/重量'), ('第 5 列', '货物总值'), ('第 6 列', '包装种类及数量')]),
                TextRow('', '第 1 页 OCR第 2 行', [('第 1 列', '测试罐头'), ('第 2 列', '2005800000(P.R/Q'), ('第 3 列', '测试市'), ('第 4 列', '60箱'), ('第 5 列', '660 美元'), ('第 6 列', '纸箱')]),
                TextRow('', '第 1 页 OCR第 3 行', [('第 4 列', '120千克'), ('第 6 列', '/箱/60')])]
        found = values(rows)
        self.assertEqual(found.get('amount'), ['660 美元'])
        self.assertEqual(found.get('declared_weight'), ['120千克'])
        self.assertEqual(found.get('currency'), ['美元'])
        self.assertEqual(found.get('quantity'), ['60箱'])

    def test_flat_customs_detail_and_multiple_inline_labels(self):
        found = values([TextRow('检验号：TEST001 批号：20261012', '段落 1'),
                        TextRow('2005800000         测试罐头       120KG       6.6 / 660.00 /美金       中国', '表格 1 第 1 行')])
        self.assertEqual(found.get('document_no'), ['TEST001'])
        self.assertEqual(found.get('batch_no'), ['20261012'])
        self.assertEqual(found.get('hs_code'), ['2005800000'])
        self.assertEqual(found.get('item_description'), ['测试罐头'])
        self.assertEqual(found.get('unit_price'), ['6.6'])
        self.assertEqual(found.get('amount'), ['660.00'])

    def test_inline_form_is_not_a_column_header(self):
        rows = [TextRow('', '工作表「INV」第 1 行', [('A1', 'INVOICE NO. 发票号:'), ('C1', 'TEST123'), ('E1', 'CONTRACT NO. 合同号:'), ('G1', 'SC123')])]
        found = values(rows)
        self.assertEqual(found.get('document_no'), ['TEST123'])
        self.assertEqual(found.get('contract_no'), ['SC123'])

    def test_inline_specification_preserves_product_description(self):
        found = values([TextRow('DESCRIPTION: TEST CORN 24x320/285gr EOE', 'A1')])
        self.assertEqual(found.get('item_description'), ['TEST CORN EOE'])

    def test_customs_currency_is_not_a_number_and_headers_not_data(self):
        self.assertNotIn('currency', values([TextRow('币制: 12.3', 'A1')]))
        self.assertNotIn('package_type', values([TextRow('包装种类及数量', 'A2')]))

    def test_short_to_label_does_not_match_total(self):
        self.assertNotIn('buyer_name', values([TextRow('TOTAL', 'A1'), TextRow('TOTAL SAY DOLLARS ONLY', 'A2')]))

    def test_ocr_separate_date_boxes_and_voyage_value(self):
        import numpy as np
        from types import SimpleNamespace
        from unittest.mock import patch
        from PIL import Image
        import server.trade_documents as module
        result = SimpleNamespace(txts=['日期：', '2026-10-12', '航次：V.2642E'],
            scores=np.array([0.99, 0.98, 0.99]),
            boxes=np.array([[[10,10],[60,10],[60,30],[10,30]], [[65,10],[180,10],[180,30],[65,30]], [[10,40],[180,40],[180,60],[10,60]]]))
        with patch.object(module, '_ocr_engine', return_value=lambda image: result):
            found = values(module._ocr_image_local(Image.new('RGB', (200, 100)), '第 1 页 OCR'))
        self.assertEqual(found.get('document_date'), ['2026-10-12'])
        self.assertEqual(found.get('voyage_no'), ['V.2642E'])

    def test_pdf_labels_and_values_follow_visual_positions(self):
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((50, 90), 'Port of discharge:')
        page.insert_text((50, 120), 'Proforma ETD:')
        page.insert_text((200, 90), 'TEST PORT')
        page.insert_text((200, 120), '2026-10-21')
        rows, _ = _read_document('test.pdf', doc.tobytes())
        doc.close()
        found = values(rows)
        self.assertEqual(found.get('destination_port'), ['TEST PORT'])
        self.assertEqual(found.get('shipment_date'), ['2026-10-21'])

    def test_word_merged_labels_and_multiline_values(self):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=2, cols=4)
        table.cell(0, 0).merge(table.cell(0, 1)).text = '产品名称'
        table.cell(0, 2).merge(table.cell(0, 3)).text = '测试罐头'
        table.cell(1, 0).text = '毛重（千克）\n120KG'
        table.cell(1, 1).text = '数量\n60箱'
        table.cell(1, 2).text = '运抵国（地区）\nTESTLAND'
        output = io.BytesIO()
        doc.save(output)
        rows, _ = _read_document('test.docx', output.getvalue())
        found = values(rows)
        self.assertEqual(found.get('item_description'), ['测试罐头'])
        self.assertEqual(found.get('gross_weight'), ['120KG'])
        self.assertEqual(found.get('quantity'), ['60箱'])
        self.assertEqual(found.get('destination_country'), ['TESTLAND'])

    def test_units_in_headers_and_multiple_product_rows(self):
        rows = [TextRow('', '工作表「PI」第 1 行', [('A1', 'DESCRIPTION'), ('B1', 'Q\'TY /20GP CTNS'), ('C1', 'UNIT PRICE\nFOB TEST\nUSD/CTN'), ('D1', 'AMOUNT\nUSD')]),
                TextRow('', '工作表「PI」第 2 行', [('A2', 'TEST CORN'), ('B2', '100'), ('C2', '6.6'), ('D2', '660')]),
                TextRow('', '工作表「PI」第 3 行', [('A3', 'TEST PEAS'), ('B3', '200'), ('C3', '7'), ('D3', '1400')])]
        found = values(rows)
        self.assertEqual(found.get('item_description'), ['TEST CORN', 'TEST PEAS'])
        self.assertEqual(found.get('quantity_per_container'), ['100', '200'])
        self.assertEqual(found.get('unit_price'), ['6.6', '7'])
        self.assertEqual(found.get('amount'), ['660', '1400'])

    def test_adjacent_headers_are_not_values(self):
        rows = [TextRow('', '工作表「PI」第 1 行', [('A1', 'DESCRIPTION'), ('B1', 'PACKING'), ('C1', 'LID TYPE')]),
                TextRow('', '工作表「PI」第 2 行', [('A2', 'TEST CORN'), ('B2', '24TINS/CTN'), ('C2', 'EASY OPEN')])]
        self.assertEqual(values(rows).get('package_type'), ['24TINS/CTN'])

    def test_declaration_number_prefix_is_not_specification(self):
        found = values([TextRow('3. 包装规格：铁罐装，纸箱装', 'A1'),
                        TextRow('0. 品牌类型：境外品牌', 'A2'),
                        TextRow('4. 品牌：TEST BRAND', 'A3')])
        self.assertEqual(found.get('item_specification'), ['铁罐装，纸箱装'])
        self.assertEqual(found.get('brand_type'), ['境外品牌'])
        self.assertEqual(found.get('brand'), ['TEST BRAND'])

    def test_empty_pdf_label_does_not_consume_next_label_or_footer_date(self):
        rows = [TextRow('Final destination:', '第 1 页'), TextRow('温馨提示：', '第 1 页'),
                TextRow('编号:TEST 版本号:1 日期: 01/08/1995 修改版本号:', '第 1 页')]
        found = values(rows)
        self.assertNotIn('destination_country', found)
        self.assertNotIn('document_date', found)


if __name__ == '__main__':
    unittest.main()
