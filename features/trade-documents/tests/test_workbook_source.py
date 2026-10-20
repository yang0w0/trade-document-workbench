"""Workbook source API regressions with synthetic files only."""
import io
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / 'application'))
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from server.trade_documents import router


class WorkbookSourceTests(unittest.TestCase):
    def setUp(self):
        app = FastAPI()
        app.include_router(router)
        self.client = TestClient(app)

    def test_xls_source_styles_focus_original_and_readonly(self):
        content = (Path(__file__).parent / 'fixtures/synthetic-source.xls').read_bytes()
        response = self.client.post('/api/documents/workbook-sessions', files={'file': ('arbitrary.XLS', content)})
        self.assertEqual(response.status_code, 200, response.text)
        base = '/api/documents/workbook-sessions/' + response.json()['sessionId']
        response = self.client.get(base, params={'sheet': '来源表', 'cell': 'B2'})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual((data['focusSheet'], data['focusCell']), ('来源表', 'B2'))
        self.assertEqual([s['name'] for s in data['sheets']], ['来源表', '第二表'])
        sheet = data['sheets'][0]
        self.assertEqual(sheet['rows'][1][1], 125)
        self.assertTrue(sheet['rows'][2][1].startswith('2026-10-19'))
        self.assertEqual(sheet['merged'][0], {'minRow': 1, 'maxRow': 1, 'minColumn': 1, 'maxColumn': 3})
        style = data['styles'][sheet['styleIds'][0][0]]
        self.assertTrue(style['font']['bold'])
        self.assertEqual(style['font']['color'], '#ff0000')
        self.assertEqual(style['fill'], '#ffff00')
        self.assertEqual(style['borders']['bottom']['style'], 'thin')
        self.assertEqual(sheet['columns'][0]['width'], 20)
        self.assertEqual(sheet['rowDimensions']['1']['height'], 25)
        self.assertFalse(data['capabilities']['edit'])
        self.assertEqual(self.client.get(base + '/original').content, content)
        self.assertEqual(self.client.post(base + '/save', json={'changes': []}).status_code, 400)

    def test_xlsx_still_edits_and_downloads_original(self):
        book = Workbook()
        book.active['B2'] = 125
        stream = io.BytesIO()
        book.save(stream)
        content = stream.getvalue()
        response = self.client.post('/api/documents/workbook-sessions', files={'file': ('invoice.xlsx', content)})
        self.assertEqual(response.status_code, 200, response.text)
        base = '/api/documents/workbook-sessions/' + response.json()['sessionId']
        self.assertTrue(self.client.get(base).json()['capabilities']['edit'])
        response = self.client.post(base + '/save', json={'changes': [{'sheet': 'Sheet', 'cell': 'B2', 'value': 200}]})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active['B2'].value, 200)
        self.assertEqual(self.client.get(base + '/original').content, content)

    def test_corrupt_xls_returns_actionable_error(self):
        response = self.client.post('/api/documents/workbook-sessions', files={'file': ('broken.xls', b'broken')})
        self.assertEqual(response.status_code, 400)
        self.assertIn('无法打开', response.json()['detail'])


if __name__ == '__main__':
    unittest.main()
